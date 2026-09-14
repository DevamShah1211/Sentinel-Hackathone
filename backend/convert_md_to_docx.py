import os
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_color):
    """Set background color of a table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_color)
    tcPr.append(shd)

def convert_md_to_docx():
    md_file = Path(r"d:\Hacathone CCTV\sentinel-platform\submission\FORM_SUBMISSION_FIELD_ANSWERS.md")
    docx_file = Path(r"d:\Hacathone CCTV\sentinel-platform\submission\FORM_SUBMISSION_FIELD_ANSWERS.docx")

    if not md_file.exists():
        print(f"File not found: {md_file}")
        return

    doc = docx.Document()

    # Set page margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_p.add_run("Sentinel — Hackathon Submission Form Answers")
    run.font.name = "Calibri"
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 32, 67) # Deep navy

    subtitle_p = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = subtitle_p.add_run("Gujarat CCTV Integration Hackathon 2026 · Category 1\nCopy-Paste Answers Cheat Sheet")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(100, 110, 120)

    doc.add_paragraph() # Spacing

    with open(md_file, "r", encoding="utf-8") as f:
        lines = f.readlines()

    current_h2 = ""
    current_h3 = ""

    for line in lines:
        line_str = line.strip()
        if not line_str:
            continue

        if line_str.startswith("# "):
            continue
        elif line_str.startswith("## "):
            h2_text = line_str.replace("## ", "").strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(h2_text)
            run.font.name = "Calibri"
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(28, 54, 110)
        elif line_str.startswith("### "):
            h3_text = line_str.replace("### ", "").strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(h3_text)
            run.font.name = "Calibri"
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 40, 60)
        elif line_str.startswith("- [x] "):
            item_text = line_str.replace("- [x] ", "").replace("**", "").strip()
            p = doc.add_paragraph(style='List Bullet')
            run = p.add_run(f"☑ {item_text}")
            run.font.name = "Calibri"
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 120, 60) # Green check
        elif line_str.startswith("- "):
            item_text = line_str.replace("- ", "").strip()
            p = doc.add_paragraph(style='List Bullet')
            run = p.add_run(item_text)
            run.font.name = "Calibri"
            run.font.size = Pt(11)
        elif line_str.startswith("http://") or line_str.startswith("https://"):
            table = doc.add_table(rows=1, cols=1)
            cell = table.cell(0, 0)
            set_cell_background(cell, "F0F4F8")
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(line_str)
            run.font.name = "Consolas"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0, 90, 180)
            doc.add_paragraph()
        elif not line_str.startswith("---") and not line_str.startswith("*"):
            table = doc.add_table(rows=1, cols=1)
            cell = table.cell(0, 0)
            set_cell_background(cell, "F8F9FA")
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(line_str)
            run.font.name = "Segoe UI"
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGBColor(20, 20, 20)
            doc.add_paragraph()

    doc.save(docx_file)
    print(f"SUCCESS: Generated Word document at {docx_file}")

if __name__ == "__main__":
    convert_md_to_docx()
